import pathlib
import re
import pymupdf  # fix fitz deprecated
import docx
import openpyxl

def _clean(s):
    if s is None: return ""
    # handle jika yang masuk adalah objek Cell, Paragraph, dll
    if not isinstance(s, str):
        s = str(s) if not hasattr(s, 'text') else s.text
    return s.strip()

def extract_text(uploaded_file):
    name = uploaded_file.name.lower()
    text = ""
    meta = {"kode_mk": None, "sks": None, "dosen": None}
    
    try:
        data = uploaded_file.read()
        uploaded_file.seek(0)

        if name.endswith(".pdf"):
            doc = pymupdf.open(stream=data, filetype="pdf")
            for p in doc:
                text += p.get_text("text") + "\n"

        elif name.endswith(".docx"):
            d = docx.Document(uploaded_file)
            # 1. Paragraf
            for p in d.paragraphs:
                text += _clean(p.text) + "\n"
            # 2. Tabel (INI YANG BIKIN ERROR TADI)
            for table in d.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        row_text.append(_clean(cell.text))
                    text += " | ".join(row_text) + "\n"

        elif name.endswith(".doc"):
            # .doc legacy, coba baca sebagai docx, jika gagal pakai decode
            try:
                d = docx.Document(uploaded_file)
                for p in d.paragraphs:
                    text += _clean(p.text) + "\n"
            except:
                text = data.decode(errors="ignore")

        elif name.endswith((".xlsx", ".xls")):
            wb = openpyxl.load_workbook(uploaded_file, data_only=True)
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    vals = [ _clean(c) for c in row if c is not None ]
                    if vals:
                        text += " | ".join(vals) + "\n"
        else:
            text = data.decode(errors="ignore")

        # Ekstraksi identitas - HARAM ngarang, kalau tidak ketemu biarkan None
        m_kode = re.search(r"\b(?:IF|TI|MK)\s*[-]?\s*\d{3,4}[A-Z]?\b", text, re.I)
        if m_kode: meta["kode_mk"] = _clean(m_kode.group(0))
        
        m_sks = re.search(r"(\d+)\s*SKS", text, re.I)
        if m_sks: meta["sks"] = _clean(m_sks.group(1))

        m_dosen = re.search(r"Dosen\s*[:]\s*(.+)", text, re.I)
        if m_dosen: meta["dosen"] = _clean(m_dosen.group(1))[:100]

        if not text.strip():
            text = "Gagal ekstrak: dokumen kosong atau hanya gambar"

    except Exception as e:
        text = f"Gagal ekstrak: {e}"
        # JANGAN throw, biarkan app tetap jalan dengan pesan Anti-Gagal §7

    return text, meta