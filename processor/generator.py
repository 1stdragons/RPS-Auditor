"""Generator OPSI C — RPS Auditor Polsel V5 (pasal 11, 12, 13 .clinerules).

Pasal 11: muat file asli sebagai template via Document(original_path) — JANGAN
Document() baru — pertahankan struktur 100% (kop, tabel, font, tanda tangan),
timpa rumusan CPMK lama dengan hasil V5, lalu tambahkan LAMPIRAN C Log Audit
Perubahan di belakang. Hasil dikembalikan sebagai BytesIO.
Pasal 12: anti-KeyError — template Polsel tidak punya 'Light Grid',
'List Number', 'Heading 2'; semua pemakaian style dibungkus try/except
KeyError dengan fallback paragraf polos, dan style tabel ditiru dari
tabel pertama dokumen (bukan hardcode).
Pasal 13: konsumen teks aman — app membatasi safe = text[:3500] dan
max_tokens=2500, master prompt < 2500 char.
"""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any

import openpyxl
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from openpyxl.styles import Alignment, Font

# Kalimat wajib §6 — JANGAN DIUBAH.
KALIMAT_VERIFIKASI = (
    "Sumber resmi tidak berhasil diverifikasi pada percobaan ini, "
    "CPL/SKKNI ditandai sebagai usulan auditor"
)

LOGO_PLACEHOLDER = "[LOGO POLSEL - HANYA IKON 3 WARNA TANPA TULISAN]"

# 7 poin log Before -> After (pasal 11) — dipakai DOCX & XLSX agar konsisten.
LOG_7_POIN: list[tuple[str, str, str]] = [
    (
        "CPMK",
        "C1-C2 dominan, rumusan umum tanpa objek",
        "CPMK baru dominan C3-C6 dengan KKO Vokasi (Menerapkan, Menganalisis, Menguji, Merancang)",
    ),
    (
        "Sub-CPMK",
        "Tanpa konteks industri",
        "Konteks sawit: Blok A1, TPH, TBS, Timbangan PKS, Sortasi Grade A/B/C, Premi",
    ),
    (
        "Metode",
        "Ceramah dominan",
        "Project Based Learning + Praktikum Flowchart Blok->TPH->Truk->Timbangan PKS->Sortasi",
    ),
    (
        "Rubrik Pertemuan 6",
        "Tanpa rubrik flowchart",
        "4 kriteria x skor 0-5 = 20: simbol, Grade/Premi, alur Blok->PKS, aturan kasus dosen",
    ),
    (
        "Kasus Uji Pertemuan 15",
        "Tanpa kasus uji",
        "3 kasus uji: Normal 25 janjang, Batas 0 janjang, Tidak Valid huruf/simbol "
        "+ Expected vs Actual, Bug log, Fix, Re-test, Video edukasi 2 menit",
    ),
    (
        "Bobot Penilaian",
        "Distribusi bobot tidak rinci",
        "Tabel 2 total 100% tanpa double counting; Minggu 8 UTS Praktik, Minggu 16 UAS Project Demo",
    ),
    ("Verifikasi Sumber", "Tanpa penandaan status sumber", KALIMAT_VERIFIKASI),
]

HEADER_16 = [
    "Minggu", "Sub-CPMK [Bloom]", "Bahan Kajian Sawit", "Metode", "Waktu",
    "Sumber", "Skenario Praktik Data Sawit", "Indikator Bloom",
    "Kriteria/Rubrik Koreksi", "Bobot",
]
HEADER_MTRIX = ["CPL", "CPMK", "Sub-CPMK"]
HEADER_PENILAIAN = ["No", "Komponen", "Bloom", "Indikator Sawit", "Instrumen", "Bobot"]


def _safe_get_style(doc: Any, preferred: list[str], fallback: str | None = None) -> str | None:
    """Kembalikan nama style yang benar-benar ada di doc.styles, atau fallback.

    Pengecekan memakai jalur lookup yang SAMA dengan add_paragraph
    (doc.styles[nama]) sehingga dijamin tidak KeyError — pasal 12.
    """
    for nama in preferred:
        if not nama:
            continue
        try:
            doc.styles[nama]  # melempar KeyError jika style tidak ada
            return nama
        except KeyError:
            continue
        except Exception:
            continue
    return fallback


def _safe_add_paragraph(doc: Any, teks: str, style: str | None = None) -> Any:
    """Tambah paragraf anti-KeyError: style opsional, selalu ada fallback polos.

    Pasal 12: doc.add_paragraph(text) tanpa style parameter aman karena tidak
    pernah melakukan lookup nama style.
    """
    if style:
        try:
            return doc.add_paragraph(teks, style=style)
        except KeyError:
            pass
    return doc.add_paragraph(teks)


def _safe_add_heading(doc: Any, teks: str, level: int = 1) -> Any:
    """Tambah heading sesuai pasal 12: doc.add_heading(text, level) TANPA
    .style.name = ...; jika style heading tidak ada di template -> paragraf polos tebal.
    """
    gaya = _safe_get_style(doc, [f"Heading {level}", f"heading {level}", f"Judul {level}"])
    if gaya:
        try:
            return doc.add_heading(teks, level)  # dilarang .style.name = ...
        except KeyError:
            pass
    p = _safe_add_paragraph(doc, teks)
    if p is not None and getattr(p, "runs", None):
        p.runs[0].bold = True
    return p


def _nama_style_tabel_asli(doc: Any) -> str | None:
    """Ambil nama style dari tabel pertama dokumen untuk ditiru (pasal 12)."""
    try:
        if doc.tables:
            st = doc.tables[0].style  # bisa KeyError jika rujukan style hilang
            if st is not None and st.name:
                return str(st.name)
    except Exception:
        return None
    return None


def _safe_add_table(doc: Any, rows: int, cols: int, clone_style: str | None = None) -> Any:
    """Tambah tabel TANPA style hardcode ('Light Grid' dilarang); clone dari
    style tabel asli bila ada, semua dibungkus try/except KeyError (pasal 12).
    """
    try:
        table = doc.add_table(rows, cols)  # style=None -> tidak lookup nama style
    except KeyError:
        return None
    if clone_style:
        try:
            table.style = clone_style
        except (KeyError, ValueError):
            pass  # biarkan ikut default template — jangan pernah crash
    return table


def _ganti_teks_paragraf(p: Any, teks: str) -> None:
    """Ganti teks paragraf sambil mempertahankan style/format paragraf asli."""
    try:
        if getattr(p, "runs", None):
            p.runs[0].text = teks
            for r in p.runs[1:]:
                r.text = ""
        else:
            p.add_run(teks)
    except Exception:
        pass  # struktur asli tetap utuh


_POLA_CPMK = re.compile(r"^(?:\d+[\.\)]\s*)?(?:Sub-)?CPMK\s*\d", re.I)


def _semua_paragraf(doc: Any):
    """Semua paragraf dokumen dalam urutan asli (termasuk di tabel, tanpa duplikat)."""
    for p_el in doc._element.iter(qn("w:p")):
        yield Paragraph(p_el, doc)


def _baris_cpmk(ai_text: str) -> list[str]:
    """Ambil baris hasil V5 yang berupa rumusan CPMK/Sub-CPMK (siap ditimpa)."""
    hasil: list[str] = []
    for b in (ai_text or "").splitlines():
        bersih = re.sub(r"\*\*|__|`", "", b).strip()
        if _POLA_CPMK.match(bersih) and len(bersih) > 15:
            hasil.append(bersih)
    return hasil


def _timpa_cpmk(doc: Any, ai_text: str) -> tuple[int, int]:
    """Timpa rumusan CPMK lama di dalam dokumen dengan baris hasil V5 (pasal 11).

    Best-effort dan mempertahankan style paragraf asli; hanya paragraf rumusan
    panjang yang diawali 'CPMK n' / 'Sub-CPMK n.n' yang ditimpa — judul bagian
    dan sel referensi pendek ('CPMK 1') dibiarkan utuh. Mengembalikan
    (jumlah_tertimpa, jumlah_baris_v5).
    """
    baris_baru = _baris_cpmk(ai_text)
    if not baris_baru:
        return 0, 0
    tertulis = 0
    for p in _semua_paragraf(doc):
        if tertulis >= len(baris_baru):
            break
        teks = (p.text or "").strip()
        if len(teks) <= 20 or not _POLA_CPMK.match(teks):
            continue
        _ganti_teks_paragraf(p, baris_baru[tertulis])
        tertulis += 1
    return tertulis, len(baris_baru)


def _pecah_paragraf(teks: str, maks: int = 1800) -> list[str]:
    """Pecah teks panjang menjadi potongan paragraf yang mudah dibaca."""
    hasil: list[str] = []
    buf = ""
    for b in (teks or "").splitlines():
        if buf and len(buf) + len(b) + 1 > maks:
            hasil.append(buf)
            buf = ""
        buf = f"{buf}\n{b}" if buf else b
    if buf:
        hasil.append(buf)
    return hasil or ([teks[:maks]] if teks else [])


def _tabel_rubrik(doc: Any, clone_style: str | None) -> None:
    """Rubrik flowchart Pertemuan 6 — 4 kriteria wajib x skor 0-5 = 20 (§5a)."""
    baris = [
        ("1. Ketepatan simbol (terminator, process, decision, I/O)", "0-5", "Simbol flowchart sesuai standar"),
        ("2. Kelengkapan kondisi Grade A/B/C dan Premi", "0-5", "Semua cabang kondisi grade/premi ada"),
        ("3. Konsistensi alur Blok->TPH->Truk->Timbangan PKS->Sortasi", "0-5", "Alur produksi sawit berurutan utuh"),
        ("4. Kesesuaian dengan aturan kasus dosen", "0-5", "Sesuai aturan yang diberikan dosen"),
    ]
    tabel = _safe_add_table(doc, len(baris) + 1, 3, clone_style)
    if tabel is None:
        _safe_add_paragraph(
            doc,
            "Rubrik 4 kriteria x skor 0-5 = total 20: simbol; Grade/Premi; "
            "alur Blok->TPH->Truk->Timbangan PKS->Sortasi; aturan kasus dosen.",
        )
        return
    for j, t in enumerate(("Kriteria", "Skor", "Deskripsi")):
        tabel.rows[0].cells[j].text = t
    for i, b in enumerate(baris, start=1):
        for j in range(3):
            tabel.rows[i].cells[j].text = b[j]


def _tabel_kasus_uji(doc: Any, clone_style: str | None) -> None:
    """3 kasus uji Pertemuan 15 & Sub-CPMK 4.2: 25 / 0 / huruf (§5b)."""
    baris = [
        ("Normal", "TBS 25 janjang", "Dihitung, premi sesuai Grade A/B/C", "Catat Actual; jika beda -> Bug log -> Fix -> Re-test -> Video 2 menit"),
        ("Batas", "0 janjang", "Ditolak/divalidasi dengan pesan jelas", "Catat Actual; jika beda -> Bug log -> Fix -> Re-test -> Video 2 menit"),
        ("Tidak Valid", "Huruf/simbol", "Ditolak, tidak memproses data", "Catat Actual; jika beda -> Bug log -> Fix -> Re-test -> Video 2 menit"),
    ]
    tabel = _safe_add_table(doc, len(baris) + 1, 4, clone_style)
    if tabel is None:
        _safe_add_paragraph(
            doc,
            "3 kasus uji: Normal TBS 25 janjang, Batas 0 janjang, Tidak Valid huruf/simbol "
            "+ Expected vs Actual, Bug log, Fix, Re-test, Video edukasi 2 menit.",
        )
        return
    for j, t in enumerate(("Kasus", "Input", "Expected", "Actual / Bug log / Fix / Re-test / Video")):
        tabel.rows[0].cells[j].text = t
    for i, b in enumerate(baris, start=1):
        for j in range(4):
            tabel.rows[i].cells[j].text = b[j]


def create_docx_preserve_and_replace(original_path: str, ai_text: str, meta: dict | None = None) -> BytesIO:
    """OPSI C (pasal 11): load file asli sebagai template, timpa isi, lampiran.

    WAJIB Document(original_path) — jangan Document() baru — supaya kop, tabel,
    font, dan tanda tangan (Penyusun, Ka Prodi, Kajur, P2M, WD1, Direktur)
    tetap 100% sama. Isi CPMK ditimpa hasil V5, lalu LAMPIRAN C Log Audit
    Perubahan ditambahkan di belakang. Mengembalikan BytesIO.
    """
    doc = Document(original_path)  # WAJIB: template asli, struktur 100% utuh
    meta = meta or {}

    # Deteksi style yang BENAR-BENAR ada di template (pasal 12 — anti KeyError).
    style_normal = _safe_get_style(doc, ["Normal"], fallback=None)
    style_tabel = _nama_style_tabel_asli(doc)

    # (1) TIMPA ISI — CPMK lama diganti CPMK baru hasil audit V5 (pasal 11).
    tertulis, jumlah_baru = _timpa_cpmk(doc, ai_text or "")

    # (2) LAMPIRAN C di belakang — add_page_break lalu isi lampiran.
    doc.add_page_break()
    _safe_add_heading(
        doc,
        "LAMPIRAN C: LOG AUDIT PERUBAHAN V5 — OPSI C (Struktur 100% Sama, Isi Ditimpa)",
        1,
    )
    p_logo = _safe_add_paragraph(doc, LOGO_PLACEHOLDER, style=style_normal)
    if p_logo is not None and getattr(p_logo, "runs", None):
        p_logo.runs[0].bold = True

    _safe_add_heading(doc, "A. Log Perubahan Before -> After (7 Poin)", 2)
    for i, (bagian, sebelum, sesudah) in enumerate(LOG_7_POIN, start=1):
        _safe_add_paragraph(doc, f"{i}. {bagian}: {sebelum} -> {sesudah}", style=style_normal)

    if tertulis:
        catatan = (
            f"Paragraf rumusan CPMK lama tertimpa: {tertulis} dari {jumlah_baru} baris hasil V5; "
            "sisa hasil lengkap tercantum pada bagian F lampiran ini."
        )
    else:
        catatan = (
            "Tidak ada paragraf CPMK lama yang dapat ditimpa otomatis — struktur asli dipertahankan "
            "100% dan seluruh hasil audit V5 tercantum pada bagian F lampiran ini."
        )
    _safe_add_paragraph(doc, catatan, style=style_normal)

    _safe_add_heading(doc, "B. Rubrik Flowchart Pertemuan 6 (4 Kriteria Wajib)", 2)
    _tabel_rubrik(doc, style_tabel)

    _safe_add_heading(doc, "C. 3 Kasus Uji Pertemuan 15 & Sub-CPMK 4.2", 2)
    _tabel_kasus_uji(doc, style_tabel)
    _safe_add_paragraph(
        doc,
        "Untuk setiap kasus wajib dicatat: Expected vs Actual, Bug log, Fix, Re-test, "
        "dan Video edukasi 2 menit.",
        style=style_normal,
    )

    _safe_add_heading(doc, "D. Verifikasi Sumber (Wajib)", 2)
    _safe_add_paragraph(doc, KALIMAT_VERIFIKASI, style=style_normal)

    _safe_add_heading(doc, "E. Identitas dari Dokumen Sumber", 2)
    for label, kunci in (("Kode MK", "kode_mk"), ("SKS", "sks"), ("Dosen Pengampu", "dosen")):
        nilai = meta.get(kunci) or "Sesuai dokumen sumber"
        _safe_add_paragraph(doc, f"{label}: {nilai}", style=style_normal)

    _safe_add_heading(doc, "F. Detail Hasil Audit V5 (Proses AI Belakang Layar)", 2)
    for potongan in _pecah_paragraf(ai_text or "", 1800):
        _safe_add_paragraph(doc, potongan, style=style_normal)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _baris_tabel(b: str) -> bool:
    s = b.strip()
    return s.startswith("|") and s.endswith("|") and s.count("|") >= 2


def _pemisah(b: str) -> bool:
    s = b.strip()
    return bool(re.fullmatch(r"[\s|:\-]+", s)) and "-" in s


def _baca_tabel(baris: list[str], i: int) -> tuple[list[str], list[list[str]], int] | None:
    """Baca satu blok tabel markdown mulai indeks i -> (header, baris, indeks_berikut)."""
    if i + 1 >= len(baris) or not _baris_tabel(baris[i]) or _pemisah(baris[i]):
        return None
    if not _pemisah(baris[i + 1]):
        return None
    header = [re.sub(r"\*\*|__|`", "", x).strip() for x in baris[i].strip().strip("|").split("|")]
    rows: list[list[str]] = []
    j = i + 2
    while j < len(baris) and _baris_tabel(baris[j]) and not _pemisah(baris[j]):
        rows.append([re.sub(r"\*\*|__|`", "", x).strip() for x in baris[j].strip().strip("|").split("|")])
        j += 1
    return header, rows, j


def parse_tables(teks: str) -> list[dict[str, Any]]:
    """Kumpulkan semua tabel markdown dari teks hasil model (dipakai validator)."""
    baris = (teks or "").splitlines()
    hasil: list[dict[str, Any]] = []
    i = 0
    while i < len(baris):
        blok = _baca_tabel(baris, i)
        if blok:
            header, rows, i = blok
            hasil.append({"header": header, "rows": rows})
        else:
            i += 1
    return hasil


def _cari_tabel(tables: list[dict], *ciri: str, tolak: tuple[str, ...] = ()) -> dict | None:
    """Cari tabel markdown berdasarkan kata kunci header; tolak = ciri larangan."""
    for tb in tables:
        h = " ".join(tb["header"]).lower()
        if all(c in h for c in ciri) and not any(x in h for x in tolak):
            return tb
    return None


def _normalisasi(header: list[str], rows: list[list[str]]) -> tuple[list[str], list[list[str]]]:
    """Samakan panjang header dan baris agar aman ditulis ke sheet."""
    n = max([len(header)] + [len(r) for r in rows]) if rows else len(header)
    n = max(n, len(header))
    h = list(header) + [f"Kolom {i + 1}" for i in range(len(header), n)]
    r = [list(x) + [""] * (n - len(x)) for x in rows]
    return h, r


def _isi_sheet(ws, tabel: dict | None, header_bawaan: list[str], catatan: str) -> None:
    """Tulis satu sheet: pakai tabel hasil AI jika ada, fallback header + catatan."""
    if tabel:
        header, rows = _normalisasi(tabel["header"], tabel["rows"])
    else:
        header = list(header_bawaan)
        rows = [[catatan] + [""] * (len(header) - 1)]
    ws.append(header)
    for r in rows:
        ws.append(r)
    _hias(ws)


def _hias(ws) -> None:
    """Format sheet: header tebal, freeze baris 1, lebar kolom & wrap otomatis."""
    for sel in ws[1]:
        sel.font = Font(bold=True)
        sel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for kol in ws.iter_cols(min_col=1, max_col=ws.max_column, min_row=1, max_row=ws.max_row):
        panjang = max((len(str(c.value)) if c.value is not None else 0) for c in kol)
        ws.column_dimensions[kol[0].column_letter].width = min(max(panjang + 2, 12), 60)
    for baris in ws.iter_rows(min_row=2):
        for sel in baris:
            sel.alignment = Alignment(wrap_text=True, vertical="top")


def create_xlsx(full_text: str, log_list: list | None = None) -> BytesIO:
    """XLSX 4 sheet (§8): 16 Pertemuan, Matriks CPL-CPMK, Penilaian, Log Pembaruan."""
    tables = parse_tables(full_text or "")
    t16 = _cari_tabel(tables, "minggu") or _cari_tabel(tables, "pertemuan")
    tmtx = _cari_tabel(tables, "cpl", "cpmk") or _cari_tabel(
        tables, "cpmk", "sub-cpmk", tolak=("minggu", "bobot", "bukti")
    )
    tpen = _cari_tabel(tables, "bobot", tolak=("minggu",))

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "16 Pertemuan"
    _isi_sheet(ws1, t16, HEADER_16, "Tabel 16 pertemuan tidak terdeteksi di hasil model")
    _isi_sheet(
        wb.create_sheet("Matriks CPL-CPMK"), tmtx, HEADER_MTRIX,
        "Matriks CPL-CPMK tidak terdeteksi di hasil model",
    )
    _isi_sheet(
        wb.create_sheet("Penilaian"), tpen, HEADER_PENILAIAN,
        "Tabel penilaian tidak terdeteksi di hasil model",
    )

    # Sheet Log Pembaruan: 7 poin Before->After (pasal 11) + tambahan opsional.
    ws_log = wb.create_sheet("Log Pembaruan")
    ws_log.append(["No", "Bagian", "Before", "After"])
    for i, (bagian, sebelum, sesudah) in enumerate(LOG_7_POIN, start=1):
        ws_log.append([i, bagian, sebelum, sesudah])
    if log_list:
        for i, item in enumerate(log_list, start=len(LOG_7_POIN) + 1):
            ws_log.append([i, str(item)[:80], "", str(item)])
    _hias(ws_log)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf



